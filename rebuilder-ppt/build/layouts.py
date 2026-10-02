# -*- coding: utf-8 -*-
from engine import *
from pptx.util import Inches, Pt

BODY_Y=2.78; BODY_H=4.2

def _frame(d,spec):
    bg=spec['bg']; p=spec['p']; sl=d.new(bg)
    spec['pname']=d.parts[p-1]['name']
    d.band(sl,p,bg); code,label=spec['sec']; d.chip(sl,p,code,label,bg); d.title(sl,spec['t'],bg)
    d.footer(sl,bg); return sl

def L_bullets(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    m=spec.get('m'); tw=7.0 if m else 12.1
    d.card(sl,0.6,BODY_Y,tw+0.1,BODY_H,bg,p)
    d.text(sl,0.8,BODY_Y+0.15,tw-0.3,BODY_H-0.3,['•  '+b for b in spec['b']] if not spec.get('plain') else spec['b'],size=(fs['body']+2 if fs['body']<=20 else fs['body']),color=st[2],accent=d.col(p) if not st[4] else 'FFD479',sp=fs['sp']+2,anchor='m')
    if m: d.media(sl,7.9,BODY_Y,4.85,BODY_H,m,bg,p)
    return sl
def L_media(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']
    cap=spec.get('cap')
    h=BODY_H-(0.65 if cap else 0)
    d.media(sl,0.6,BODY_Y,12.13,h,spec['m'],bg,p)
    if cap: d.text(sl,0.6,BODY_Y+h+0.1,12.13,0.55,[cap],size=d.fs['cap'],color=st[2],bold=True,align='c',anchor='m',accent=d.col(p) if not st[4] else 'FFD479')
    return sl
def L_compare(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    (lh,li),(rh,ri)=spec['l'],spec['r']; cols=spec.get('cols',(d.col(p),'7A8794'))
    w=5.85
    for i,(h,items) in enumerate([(lh,li),(rh,ri)]):
        x=0.6+i*(w+0.43)
        d.card(sl,x,BODY_Y,w,BODY_H,bg,0) or None
        d.rect(sl,x,BODY_Y,w,0.62,fill=cols[i],rad=0.14)
        d.text(sl,x,BODY_Y,w,0.62,[h],size=fs['body']+2,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
        d.text(sl,x+0.2,BODY_Y+0.75,w-0.4,BODY_H-0.85,['•  '+t for t in items],size=(fs['body']+2 if fs['body']<=20 else fs['body']-1),color=st[2],sp=fs['sp']+4,accent=cols[i] if not st[4] else 'FFD479',anchor='m')
    if spec.get('mid'):
        d.oval(sl,6.37,BODY_Y+BODY_H/2-0.3,0.6,0.6,'FFFFFF'); d.text(sl,6.37,BODY_Y+BODY_H/2-0.3,0.6,0.6,[spec['mid']],size=14,color='333333',bold=True,align='c',anchor='m',sp=0)
    return sl
def L_cards(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    cs=spec['c']; n=len(cs); foot=spec.get('foot'); H=BODY_H-(0.8 if foot else 0)
    gap=0.25; w=(12.13-gap*(n-1))/n
    for i,c in enumerate(cs):
        x=0.6+i*(w+gap); col=d.pal[(p-1+i)%len(d.pal)] if spec.get('multi') else d.col(p)
        d.card(sl,x,BODY_Y,w,H,bg,0)
        d.oval(sl,x+0.2,BODY_Y+0.2,0.55,0.55,col); d.text(sl,x+0.2,BODY_Y+0.2,0.55,0.55,[str(i+1)],size=18,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
        Lm=max(sum(1.0 if ord(ch)>0x2000 else 0.55 for ch in cc[0].replace('**','')) for cc in cs)
        ts=int(max(fs['body']-4,min(fs['body']+3,(w-0.3)*72*1.8/(Lm*1.1))))
        d.text(sl,x+0.12,BODY_Y+0.85,w-0.24,1.0,[c[0]],size=ts,color=st[2],bold=True,sp=0,anchor='m')
        if len(c)>1: d.text(sl,x+0.12,BODY_Y+1.95,w-0.24,H-2.05,[c[1]],size=max(13,ts-3),color=st[3] if st[4] else '4B5560',sp=0)
    if foot:
        d.rect(sl,0.6,BODY_Y+H+0.12,12.13,0.62,fill=d.col(p),rad=0.14)
        d.text(sl,0.6,BODY_Y+H+0.12,12.13,0.62,[foot],size=fs['body']-2,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
    return sl
def L_timeline(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    ev=spec['ev']; n=len(ev); d.card(sl,0.6,BODY_Y,12.13,BODY_H,bg,0)
    y=BODY_Y+1.35; d.rect(sl,1.0,y,11.33,0.08,fill=d.col(p),rad=0.04)
    w=11.33/n
    for i,(dt,lb) in enumerate(ev):
        cx=1.0+w*i+w/2; d.oval(sl,cx-0.2,y-0.16,0.4,0.4,d.col(p))
        d.text(sl,cx-w/2+0.04,BODY_Y+0.35,w-0.08,0.7,[dt],size=fs['body'],color=d.col(p) if not st[4] else 'FFD479',bold=True,align='c',anchor='b',sp=0)
        d.text(sl,cx-w/2+0.06,y+0.4,w-0.12,BODY_H-1.95,[lb],size=fs['body']-3,color=st[2],align='c',sp=0)
    if spec.get('foot'): d.text(sl,0.8,BODY_Y+BODY_H-0.62,11.7,0.55,[spec['foot']],size=fs['body']-4,color=st[3] if st[4] else '4B5560',align='c',anchor='m',sp=0,bold=True,accent=d.col(p))
    return sl
def L_question(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    d.card(sl,0.6,BODY_Y,12.13,BODY_H,bg,p)
    tg=spec.get('tag','생각해요'); tw_=0.6+len(tg)*0.25
    d.rect(sl,0.9,BODY_Y+0.25,tw_,0.5,fill=d.col(p),rad=0.2); d.text(sl,0.9,BODY_Y+0.25,tw_,0.5,[tg],size=16,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
    d.text(sl,0.9,BODY_Y+0.9,11.5,1.75,[spec['q']],size=fs['q'],color=st[2],bold=True,align='c',anchor='m',sp=0,accent=d.col(p) if not st[4] else 'FFD479')
    if spec.get('hint'): d.text(sl,0.9,BODY_Y+2.8,11.5,1.25,spec['hint'],size=fs['body']-4,color=st[3] if st[4] else '4B5560',align='c',anchor='m',sp=4)
    return sl
def L_table(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    hd=spec['hd']; rows=spec['rows']; cw=spec.get('cw'); n=len(hd); 
    x,y,w,h=0.6,BODY_Y,12.13,BODY_H-0.05
    tb=sl.shapes.add_table(len(rows)+1,n,Inches(x),Inches(y),Inches(w),Inches(h)).table
    tot=sum(cw) if cw else n
    for j in range(n): tb.columns[j].width=Inches(w*((cw[j] if cw else 1)/tot))
    rh=h/(len(rows)+1)
    for i in range(len(rows)+1): tb.rows[i].height=Inches(rh)
    dark=st[4]
    for i in range(len(rows)+1):
        for j in range(n):
            c=tb.cell(i,j); c.margin_left=c.margin_right=Inches(0.1); c.margin_top=c.margin_bottom=Inches(0.03); c.vertical_anchor=MSO_ANCHOR.MIDDLE
            txt=hd[j] if i==0 else rows[i-1][j]
            tf=c.text_frame; tf.word_wrap=True; pr=tf.paragraphs[0]; pr.alignment=PP_ALIGN.CENTER if (i==0 or j>0) else PP_ALIGN.LEFT
            sz=max(fs['body']-5,15) if i else max(fs['body']-3,16)
            d._runs(pr,txt,sz,'FFFFFF' if i==0 else st[2],i==0 or j==0,'FFD479' if dark else d.col(p))
            c.fill.solid()
            if i==0: c.fill.fore_color.rgb=rgb(d.col(p) if j!=n-1 or not spec.get('lastcol') else spec['lastcol'])
            else: c.fill.fore_color.rgb=rgb(('22324A' if i%2 else '1B2A40') if dark else ('FFFFFF' if i%2 else 'F1F5F9'))
        
    return sl
def L_quote(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    d.card(sl,0.6,BODY_Y,12.13,BODY_H,bg,p)
    d.text(sl,0.9,BODY_Y+0.1,2,1.2,['“'],size=88,color=d.col(p),bold=True,sp=0)
    d.text(sl,1.3,BODY_Y+0.55,10.7,2.5,[spec['q']],size=fs['q'],color=st[2],bold=True,anchor='m',sp=0,accent=d.col(p) if not st[4] else 'FFD479')
    if spec.get('ref'): d.text(sl,1.3,BODY_Y+3.2,10.7,0.7,[spec['ref']],size=fs['body']-3,color=st[3] if st[4] else '4B5560',align='r',sp=0)
    return sl
def L_activity(d,spec):
    sl=_frame(d,spec); bg=spec['bg']; st=BGS[bg]; p=spec['p']; fs=d.fs
    m=spec.get('m'); tw=7.1 if m else 12.1
    d.card(sl,0.6,BODY_Y,tw,BODY_H,bg,p)
    d.rect(sl,0.85,BODY_Y+0.2,1.3,0.46,fill=d.col(p),rad=0.18); d.text(sl,0.85,BODY_Y+0.2,1.3,0.46,['활동'],size=16,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
    if spec.get('time'): d.text(sl,2.25,BODY_Y+0.2,3,0.46,[spec['time']],size=14,color=st[3] if st[4] else '4B5560',anchor='m',sp=0)
    steps=[f"**{i+1}.** {t}" for i,t in enumerate(spec['steps'])]
    d.text(sl,0.85,BODY_Y+0.8,tw-0.5,BODY_H-0.95,steps,size=fs['body']-2,color=st[2],sp=fs['sp'],accent=d.col(p) if not st[4] else 'FFD479',anchor='m')
    if m: d.media(sl,7.9,BODY_Y,4.83,BODY_H,m,bg,p)
    return sl
LAY={'bullets':L_bullets,'media':L_media,'compare':L_compare,'cards':L_cards,'timeline':L_timeline,'question':L_question,'table':L_table,'quote':L_quote,'activity':L_activity}

# ------------------------------------------------------------ special slides
def S_cover(d,spec):
    bg=spec['bg']; sl=d.new(bg); st=BGS[bg]; spec['pname']='-'
    dark=st[4]; c=d.pal[0]
    d.rect(sl,0.9,1.5,8.2,4.6,fill='FFFFFF' if not dark else '0B1426',a=92 if not dark else 78,rad=0.3)
    d.rect(sl,1.25,1.85,2.9,0.5,fill=c,rad=0.2); d.text(sl,1.25,1.85,2.9,0.5,['REBUILDER 통일교육'],size=15,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
    d.text(sl,1.2,2.55,7.6,2.3,[spec['t']],size=spec['size'],color=st[2] if not dark else 'FFFFFF',bold=True,anchor='m',sp=0,accent=c if not dark else 'FFD479',ls=1.05)
    d.text(sl,1.2,4.9,7.6,0.9,[spec['sub']],size=20,color=st[3] if not dark else st[3],sp=0)
    d.rect(sl,9.45,1.5,3.0,4.6,fill=c,rad=0.3)
    d.text(sl,9.45,1.85,3.0,0.6,[d.grade],size=22,color='FFFFFF',bold=True,align='c',sp=0)
    d.text(sl,9.55,2.7,2.8,3.2,spec['side'],size=16,color='FFFFFF',align='c',sp=8)
    d.text(sl,0.9,6.55,11.5,0.4,['제안: SEE-IT  |  자유와 사람을 위한 통일 — 구속사적 관점으로 배우는 한반도와 우리의 사명'],size=12,color=st[3])
    d.notes(sl,spec,'표지'); return sl
def S_intro(d,spec):
    bg=spec['bg']; sl=d.new(bg); st=BGS[bg]; spec['pname']='-'; c=d.pal[0]
    d.rect(sl,0,0,13.333,0.95,fill=c); d.text(sl,0.6,0.1,9,0.75,['강사 소개'],size=d.fs['band']+8,color='FFFFFF',bold=True,anchor='m',sp=0)
    d.card(sl,1.0,1.5,11.3,4.7,bg,0)
    lines=spec['lines']; m=spec.get('m')
    if m: d.media(sl,8.4,1.9,3.6,3.9,m,bg,0)
    d.text(sl,1.6,1.8,6.7 if m else 10.2,4.1,lines,size=d.fs['body']+2,color='FFFFFF' if st[4] else st[2],sp=14,anchor='m',accent='FFD479' if st[4] else c)
    d.footer(sl,bg); d.notes(sl,spec,'도입'); return sl
def S_goal(d,spec):
    bg=spec['bg']; sl=d.new(bg); st=BGS[bg]; spec['pname']='-'; c=d.pal[0]
    d.rect(sl,0,0,13.333,0.95,fill=c); d.text(sl,0.6,0.1,9,0.75,[spec.get('t','이번 시간에는?')],size=d.fs['band']+8,color='FFFFFF',bold=True,anchor='m',sp=0)
    d.card(sl,1.0,1.4,11.3,4.9,bg,0)
    d.text(sl,1.4,1.5,10.5,1.5,[spec['q']],size=int(d.fs['q']*0.8),color='FFFFFF' if st[4] else st[2],bold=True,align='c',anchor='m',sp=0,accent='FFD479' if st[4] else c)
    for i,g in enumerate(spec['goals']):
        y=3.3+i*0.95; d.oval(sl,1.5,y+0.08,0.6,0.6,d.col(i+1)); d.text(sl,1.5,y+0.08,0.6,0.6,[str(i+1)],size=20,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
        d.text(sl,2.3,y,9.6,0.78,[g],size=d.fs['body']-1,color='FFFFFF' if st[4] else st[2],anchor='m',sp=0,accent='FFD479' if st[4] else c)
    d.footer(sl,bg); d.notes(sl,spec,'도입'); return sl
def S_agenda(d,spec):
    bg=spec['bg']; sl=d.new(bg); st=BGS[bg]; spec['pname']='-'; c=d.pal[0]
    d.rect(sl,0,0,13.333,0.95,fill=c); d.text(sl,0.6,0.1,9,0.75,['목차 — 오늘의 대주제'],size=d.fs['band']+8,color='FFFFFF',bold=True,anchor='m',sp=0)
    n=len(d.parts); H=5.35; rh=min(1.2,H/n)
    d.card(sl,0.9,1.25,11.55,H+0.3,bg,0)
    for i,pt in enumerate(d.parts):
        y=1.45+i*rh; col=d.col(i+1)
        d.rect(sl,1.2,y+0.05,0.85,rh-0.12,fill=col,rad=0.16); d.text(sl,1.2,y+0.05,0.85,rh-0.12,[ROMAN[i]],size=24,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
        d.text(sl,2.25,y,6.0,rh-0.04,[pt['name']],size=d.fs['body']-1,color='FFFFFF' if st[4] else st[2],bold=True,anchor='m',sp=0)
        d.text(sl,8.3,y,2.7,rh-0.04,[' · '.join(c for c,_ in pt['secs'])],size=12,color=st[3],anchor='m',align='r',sp=0)
        d.text(sl,11.0,y,1.3,rh-0.04,[f"{pt['min']}분"],size=d.fs['body']-6,color=col if not st[4] else 'FFD479',bold=True,anchor='m',align='r',sp=0)
    d.footer(sl,bg); d.notes(sl,spec,'도입'); return sl
def S_divider(d,spec):
    p=spec['p']; pt=d.parts[p-1]; c=d.col(p); sl=d.new('CLEAN'); spec['bg']='DIV'; spec['pname']=pt['name']
    d.rect(sl,0,0,13.333,7.5,fill=c)
    d.rect(sl,0,0,13.333,7.5,fill='000000',a=10)
    d.text(sl,-0.3,0.2,8,7.2,[str(p)],size=420,color='FFFFFF',bold=True,sp=0,shp=None).text_frame.paragraphs[0].runs[0].font.color.rgb=rgb('FFFFFF')
    # ghost numeral 투명도
    g=sl.shapes[-1]; rp=g.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr(); sf=rp.find(qn('a:solidFill')); etree.SubElement(sf[0],qn('a:alpha')).set('val','16000')
    d.text(sl,0.8,0.55,6,0.5,[f"대주제  PART {p} / {len(d.parts)}"],size=18,color='FFFFFF',bold=True,sp=0)
    d.text(sl,0.8,2.05,6.5,3.0,[pt['name']],size=d.fs['div'],color='FFFFFF',bold=True,anchor='m',sp=0,ls=1.05)
    d.text(sl,0.8,5.3,6.4,1.4,[pt['q']],size=d.fs['body']-2,color='FFFFFF',sp=0)
    d.rect(sl,7.5,1.0,5.3,5.5,fill='FFFFFF',a=96,rad=0.25)
    d.text(sl,7.8,1.15,4.8,0.55,['중주제'],size=18,color=c,bold=True,sp=0)
    secs=pt['secs']; rh=min(1.2,4.2/len(secs))
    for i,(code,lb) in enumerate(secs):
        y=1.85+i*rh; d.rect(sl,7.8,y+0.1,0.95,rh-0.25,fill=c,rad=0.15); d.text(sl,7.8,y+0.1,0.95,rh-0.25,[code],size=18,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
        d.text(sl,8.85,y,3.8,rh-0.05,[lb],size=d.fs['body']-3,color='1F2A37',bold=True,anchor='m',sp=0)
    d.text(sl,7.8,6.0,4.8,0.4,[f"⏱ 약 {pt['min']}분"],size=14,color='5B6773',sp=0)
    spec['min']=pt['min']; spec['t']=pt['name']
    d.notes(sl,{**spec,'n':spec.get('n') or pt['q'],'bg':'CLEAN'},f"(대주제 전환) {pt['name']}")
    d.rows[-1]['bg']='DIV(파트색 단색)'; return sl
def S_thanks(d,spec):
    bg=spec['bg']; sl=d.new(bg); st=BGS[bg]; spec['pname']='-'; c=d.pal[0]; dark=st[4]
    d.rect(sl,1.5,1.6,10.3,4.0,fill='FFFFFF' if not dark else '0B1426',a=92 if not dark else 74,rad=0.35)
    d.text(sl,1.8,1.9,9.7,1.9,[spec['t']],size=d.fs['thanks'],color=st[2] if not dark else 'FFFFFF',bold=True,align='c',anchor='m',sp=0,accent=c if not dark else 'FFD479')
    d.text(sl,1.8,3.9,9.7,1.4,spec['sub'],size=d.fs['body']-2,color=st[3],align='c',anchor='m',sp=6)
    d.footer(sl,bg); d.notes(sl,spec,'마무리'); return sl
SP={'cover':S_cover,'intro':S_intro,'goal':S_goal,'agenda':S_agenda,'divider':S_divider,'thanks':S_thanks}

def build(d,specs,outdir):
    for sp in specs:
        k=sp['k']
        if k in SP: SP[k](d,sp); continue
        sl=LAY[k](d,sp); sec=f"{sp['sec'][0]} {sp['sec'][1]}"
        d.notes(sl,sp,sec)
    path=os.path.join(outdir,d.name+'.pptx'); d.prs.save(path); return path

def fit_minutes(S,PARTS):
    """슬라이드별 분(分)을 대주제 예산(PARTS[].min)에 맞춰 비례 배분(최소 1분)."""
    for p,pt in enumerate(PARTS,1):
        sl=[s for s in S if s.get('p')==p and s['k']!='divider']
        raw=[max(1,s.get('min',1)) for s in sl]; tot=sum(raw); tgt=pt['min']
        q=[r*tgt/tot for r in raw]; fl=[max(1,int(x)) for x in q]
        while sum(fl)<tgt:
            i=max(range(len(q)),key=lambda j:q[j]-fl[j]); fl[i]+=1
        while sum(fl)>tgt and max(fl)>1:
            i=max(range(len(fl)),key=lambda j:fl[j]-q[j]); fl[i]-=1
        for s,m in zip(sl,fl): s['min']=m
