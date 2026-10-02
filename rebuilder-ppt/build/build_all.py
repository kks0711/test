import sys, importlib; sys.path.insert(0,'.')
from engine import *; from layouts import build
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
D=[('deck1_chojeo','REBUILDER_통일교육_초등저학년','초등 저학년','35~40분'),('deck2_chogo','REBUILDER_통일교육_초등고학년','초등 고학년','45~50분'),('deck3_jung','REBUILDER_통일교육_중학생','중학생','45~50분'),('deck4_go','REBUILDER_통일교육_고등학생','고등학생','80~90분')]
wb=openpyxl.Workbook(); ws0=wb.active; ws0.title='배경·미디어 코드'
ws0.append(['구분','코드','설명 / 사용 규칙']); 
for c in ws0[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3A5F')
from engine import BGWHY
for k,v in BGWHY.items(): ws0.append(['배경',k,v])
ws0.append(['배경','DIV','대주제 전환 장: 파트 고유색 단색 풀블리드 + 거대 숫자 → 대/중 구조를 한눈에 구분'])
for r in [('미디어','사진','공공누리(KOGL)·NASA·UN 등 이용허락이 확인된 이미지만. 출처 캡션 필수'),('미디어','지도','위키미디어 공용 CC 지도 또는 직접 제작. 군사분계선·38선은 색·선종으로 구분'),('미디어','영상','통일부 통일교육원·북한인권정보센터 공개 자료. 상영 전 정서 안내, 3~4분 이내'),('미디어','일러스트','저작권 무료 또는 직접 제작. 국기·군복 상징 없이 평범한 복장'),('미디어','인포그래픽','구조·인과 도식. 아이콘 중심, 실존 인물 사진 사용 지양'),('미디어','활동','학생이 직접 쓰고 붙이고 말하는 참여형. 교구·활동지 동반')]: ws0.append(list(r))
ws0.column_dimensions['A'].width=8; ws0.column_dimensions['B'].width=12; ws0.column_dimensions['C'].width=110
summary=[]
for mod,name,grade,dur in D:
    WARN.clear(); m=importlib.import_module(mod)
    d=Deck(name,grade,m.PARTS,m.FS,m.PAL); p=build(d,m.S,'../decks')
    ws=wb.create_sheet(grade); ws.append(['번호','대주제','중주제','슬라이드 제목','유형','배경','미디어 유형','미디어 지시 / 준비','분']);
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3A5F')
    for r in d.rows: ws.append([r['no'],r['part'],r['sec'],r['title'].replace('**',''),r['kind'],r['bg'],r['mtype'],r['mdesc'],r['min']])
    for col,w in zip('ABCDEFGHI',[6,26,26,40,10,14,10,90,6]): ws.column_dimensions[col].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'
    summary.append((grade,dur,d.n,len(m.PARTS),[w for w in WARN if '“' not in w]))
wb.save('../REBUILDER_슬라이드별_배경·미디어_지정표.xlsx')
for s in summary: print(s)
