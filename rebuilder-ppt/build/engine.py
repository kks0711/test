# -*- coding: utf-8 -*-
"""REBUILDER 강의 PPT 엔진: 대주제(Part) 장 전환 + 중주제(Section) 칩 + 슬라이드 메시지(소) 3단 구조."""
import os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

HERE=os.path.dirname(__file__); BG=os.path.join(HERE,'..','assets','bg')
FONT="Malgun Gothic"
ROMAN=["Ⅰ","Ⅱ","Ⅲ","Ⅳ","Ⅴ","Ⅵ","Ⅶ","Ⅷ"]
WARN=[]

def rgb(h): h=h.lstrip('#'); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))

# 배경 코드별 스타일: (카드색|None, 카드투명도%, 본문색, 보조색, 다크여부)
BGS={
 'SKY':('FFFFFF',92,'1F2A37','5B6773',False),'BOARD':(None,0,'FFFFFF','CFE3D6',True),
 'NOTE':('FFFFFF',94,'1F2A37','5B6773',False),'NIGHT':('0F1B3A',70,'FFFFFF','C9D3F0',True),
 'WARM':('FFFFFF',92,'3A2A20','7A6253',False),'CLEAN':('FFFFFF',100,'1B2430','5B6773',False),
 'MAP':('FFFFFF',90,'1B2430','4F5B69',False),'PAPER':(None,0,'27231C','6A6354',False),
 'DUSK':('15163A',72,'FFFFFF','C8C6EA',True),'DAWN':('0C1A38',66,'FFFFFF','F3D9B0',True),
 'SLATE':('142033',86,'EAF0F7','9FB0C4',True),'CONTOUR':('101B2C',84,'EAF0F7','9FB0C4',True)}
BGWHY={
 'SKY':'하늘·초원 일러스트: 안전하고 밝은 첫인상, 저학년의 주의 집중과 "열린 땅" 이미지',
 'BOARD':'칠판+나무 책상: 교실 문법(질문·활동·소개)에 즉시 몰입. 흰 글씨 큰 판서',
 'NOTE':'격자 노트 질감: 생각 쓰기·정리·활동 슬라이드. 카드 위 글씨 가독성 최상',
 'NIGHT':'남색 밤하늘: 야경·단절·그리움 등 무게 있는 장면. 사진/위성이미지가 돋보임',
 'WARM':'따뜻한 석양 톤: 소망·다짐·꿈. 정서적 마무리',
 'CLEAN':'밝은 모눈 바탕: 개념·비교·정보 중심. 시각 소음 최소',
 'MAP':'옅은 등고선 지도 질감: 분단·지리·국제정치 장면에서 "한반도" 맥락 유지',
 'PAPER':'종이 질감 단색: 인용·성찰·긴 문장. 책 읽는 호흡',
 'DUSK':'짙은 남보라: 사람의 삶·인권 등 감정이 실리는 장면. 사진 몰입도 상승',
 'DAWN':'남색→새벽빛 그라데이션: 소명·미래·결단. 마지막 파트 전용',
 'SLATE':'차콜 네이비 모눈: 분석·데이터·구조도. 고등 기본 배경',
 'CONTOUR':'어두운 등고선: 지정학·안보·경로 분석. 지도·화살표 도식이 선명'}

class Deck:
    def __init__(s,name,grade_label,parts,fs,palette,part_min=None,accent_dark=None):
        s.prs=Presentation(); s.prs.slide_width=Inches(13.333); s.prs.slide_height=Inches(7.5)
        s.name=name; s.grade=grade_label; s.parts=parts; s.fs=fs; s.pal=palette
        s.rows=[]; s.n=0; s.blank=s.prs.slide_layouts[6]
    # ---------- primitives
    def _alpha(s,shp,a):
        sf=shp._element.spPr.find(qn('a:solidFill'))
        if sf is not None and a<100:
            c=sf[0]; e=etree.SubElement(c,qn('a:alpha')); e.set('val',str(int(a*1000)))
    def rect(s,sl,x,y,w,h,fill=None,a=100,line=None,lw=1,rad=None,dash=False,la=100):
        shp=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        if rad: shp.adjustments[0]=min(0.5,rad/min(w,h))
        if fill: shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); s._alpha(shp,a)
        else: shp.fill.background()
        if line:
            shp.line.color.rgb=rgb(line); shp.line.width=Pt(lw)
            if dash: shp.line.dash_style=MSO_LINE.DASH
            if la<100:
                sf=shp._element.spPr.find(qn('a:ln')).find(qn('a:solidFill')); e=etree.SubElement(sf[0],qn('a:alpha')); e.set('val',str(la*1000))
        else: shp.line.fill.background()
        shp.shadow.inherit=False
        return shp
    def oval(s,sl,x,y,w,h,fill,a=100):
        shp=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); s._alpha(shp,a); shp.line.fill.background(); shp.shadow.inherit=False; return shp
    def _runs(s,p,text,size,color,bold,accent):
        parts=re.split(r'(\*\*.+?\*\*)',text)
        for part in parts:
            if not part: continue
            b=part.startswith('**') and part.endswith('**'); t=part[2:-2] if b else part
            r=p.add_run(); r.text=t; f=r.font; f.size=Pt(size); f.bold=bool(bold or b)
            f.color.rgb=rgb(accent if b and accent else color); f.name=FONT
            rPr=r._r.get_or_add_rPr(); ea=etree.SubElement(rPr,qn('a:ea')); ea.set('typeface',FONT)
    def text(s,sl,x,y,w,h,paras,size=20,color='000000',bold=False,align='l',anchor='t',accent=None,sp=6,ls=1.12,name=None,shp=None,wrap=True):
        tb=shp or sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=tb.text_frame; tf.word_wrap=wrap
        tf.margin_left=tf.margin_right=Inches(0.08); tf.margin_top=tf.margin_bottom=Inches(0.04)
        tf.vertical_anchor={'t':MSO_ANCHOR.TOP,'m':MSO_ANCHOR.MIDDLE,'b':MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras,str): paras=[paras]
        for i,t in enumerate(paras):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}[align]
            p.space_after=Pt(sp); p.line_spacing=ls
            s._runs(p,t,size,color,bold,accent)
        s._chk(w,h,paras,size,sp,ls,name or (paras[0][:14] if paras else ''))
        return tb
    def _chk(s,w,h,paras,size,sp,ls,name):
        tot=0
        for t in paras:
            t=t.replace('**','');em=sum(1.0 if ord(c)>0x2000 else 0.52 for c in t)
            per=max(1,(w-0.2)*72/size); lines=max(1,-(-em//per)); tot+=lines*size*1.2*ls+sp
        need=tot/72+0.08
        if need>h*1.08+0.12 and size>10 and 'REBUILDER' not in name and not name.isdigit() and size<200: WARN.append(f"[{s.name}#{s.n}] 넘침 위험 '{name}' 필요 {need:.2f}in > {h:.2f}in")
    # ---------- slide frame
    def new(s,bg,part=0):
        sl=s.prs.slides.add_slide(s.blank); s.n+=1
        sl.shapes.add_picture(os.path.join(BG,bg+'.jpg'),0,0,s.prs.slide_width,s.prs.slide_height)
        return sl
    def col(s,part): return s.pal[(part-1)%len(s.pal)] if part else s.pal[0]
    def style(s,bg): return BGS[bg]
    def footer(s,sl,bg):
        st=BGS[bg]; c=st[3]
        s.text(sl,0.5,7.08,6,0.3,f"REBUILDER 통일교육  ·  {s.grade}",size=10,color=c)
        s.text(sl,11.8,7.08,1.0,0.3,str(s.n),size=10,color=c,align='r')
    def band(s,sl,part,bg):
        c=s.col(part); s.rect(sl,0,0,13.333,0.95,fill=c)
        s.rect(sl,0,0.95,13.333,0.06,fill='FFFFFF',a=35)
        pname=s.parts[part-1]['name']
        s.text(sl,0.5,0.12,1.2,0.72,ROMAN[part-1],size=26,color='FFFFFF',bold=True,anchor='m',sp=0,wrap=False)
        s.text(sl,1.75,0.12,7.2,0.72,[f"대주제 · {pname}"],size=s.fs['band'],color='FFFFFF',bold=True,anchor='m',sp=0)
        n=len(s.parts); w=0.46; gap=0.08; x0=13.333-0.5-n*(w+gap)+gap
        for i in range(n):
            x=x0+i*(w+gap); cur=(i+1==part)
            r=s.rect(sl,x,0.30,w,0.36,fill='FFFFFF' if cur else None,line='FFFFFF',lw=1.2,rad=0.12,la=100 if cur else 70)
            s.text(sl,x,0.28,w,0.4,ROMAN[i],size=13,color=c if cur else 'FFFFFF',bold=True,align='c',anchor='m',sp=0)
    def chip(s,sl,part,code,label,bg):
        st=BGS[bg]; c=s.col(part); dark=st[4]
        lw=sum(1.0 if ord(ch)>0x2000 else 0.55 for ch in label)*s.fs['chip']/72+0.5
        w=0.95+lw
        s.rect(sl,0.6,1.2,w,0.44,fill=c,rad=0.2)
        s.text(sl,0.6,1.2,0.8,0.44,code,size=s.fs['chip'],color='FFFFFF',bold=True,anchor='m',align='c',sp=0)
        s.rect(sl,1.38,1.28,0.02,0.28,fill='FFFFFF',a=70)
        s.text(sl,1.42,1.2,lw,0.44,label,size=s.fs['chip'],color='FFFFFF',bold=True,anchor='m',sp=0,wrap=False)
    def title(s,sl,t,bg,y=1.72,h=0.95):
        st=BGS[bg]; sz=s.fs['title']; 
        em=sum(1.0 if ord(c)>0x2000 else 0.55 for c in t)
        if em*sz/72>11.9: sz=int(sz*0.82)
        s.text(sl,0.55,y,12.2,h,[t],size=sz,color=st[2],bold=True,anchor='m',sp=0,name='title:'+t)
    def card(s,sl,x,y,w,h,bg,part=0,tint=False):
        st=BGS[bg]
        if st[0] is None: return None
        return s.rect(sl,x,y,w,h,fill=st[0],a=st[1],rad=0.18,line=s.col(part) if part else None,lw=1.5,la=60)
    def media(s,sl,x,y,w,h,m,bg,part=0):
        st=BGS[bg]; dark=st[4]; c=s.col(part) if part else '888888'
        s.rect(sl,x,y,w,h,fill='FFFFFF' if not dark else 'FFFFFF',a=14 if dark else 55,line=c if not dark else 'FFFFFF',lw=1.75,rad=0.12,dash=True,la=100 if not dark else 75)
        s.rect(sl,x+0.15,y+0.15,1.15,0.36,fill=c,rad=0.12)
        s.text(sl,x+0.15,y+0.15,1.15,0.36,m[0],size=12,color='FFFFFF',bold=True,align='c',anchor='m',sp=0)
        s.text(sl,x+0.3,y+0.62,w-0.6,h-0.8,[m[1]],size=15 if w>4 else 13,color=st[3] if dark else '4B5560',align='c',anchor='m',sp=0,name='media:'+m[1][:10])
    def notes(s,sl,spec,sec_label):
        m=spec.get('m'); bg=spec['bg']; L=[]
        L.append(f"【위치】대주제 {spec.get('pname','-')} ▸ 중주제 {sec_label}")
        L.append(f"【배경】{bg} — {BGWHY[bg]}")
        if m: L.append(f"【미디어】[{m[0]}] {m[1]}" + (f"\n  ▸ 준비: {m[2]}" if len(m)>2 else ''))
        if spec.get('n'): L.append(f"【진행·멘트】{spec['n']}")
        if spec.get('min'): L.append(f"【시간】약 {spec['min']}분")
        sl.notes_slide.notes_text_frame.text="\n".join(L)
        s.rows.append(dict(no=s.n,part=spec.get('pname','-'),sec=sec_label,title=spec.get('t',spec['k']),kind=spec['k'],bg=bg,mtype=(m[0] if m else '-'),mdesc=(m[1]+(' / '+m[2] if m and len(m)>2 else '') if m else '-'),min=spec.get('min','')))
